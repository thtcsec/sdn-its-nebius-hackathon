#!/usr/bin/env python
"""Build an English IEDE project report (.docx) with academic formatting + screenshots."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
IMG = ROOT / "assets" / "images" / "report"
OUT = ROOT / "docs" / "IEDE_Project_Report_Prototype.docx"


def set_run_font(run, name="Times New Roman", size=12, bold=False, italic=False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic


def add_page_number(paragraph):
    """Insert PAGE field centered in footer."""
    run = paragraph.add_run()
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char_begin)
    run._r.append(instr)
    run._r.append(fld_char_end)


def add_toc(paragraph):
    """Insert a TOC field (Word: right-click → Update Field)."""
    run = paragraph.add_run()
    fld_char_begin = OxmlElement("w:fldChar")
    fld_char_begin.set(qn("w:fldCharType"), "begin")

    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = ' TOC \\o "1-3" \\h \\z \\u '

    fld_char_separate = OxmlElement("w:fldChar")
    fld_char_separate.set(qn("w:fldCharType"), "separate")

    fld_char_end = OxmlElement("w:fldChar")
    fld_char_end.set(qn("w:fldCharType"), "end")

    run._r.append(fld_char_begin)
    run._r.append(instr)
    run._r.append(fld_char_separate)
    run._r.append(fld_char_end)


def style_doc(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(2.5)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    pf = normal.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.space_after = Pt(6)
    pf.first_line_indent = Cm(0.75)

    for level, size in ((1, 16), (2, 14), (3, 12)):
        style = styles[f"Heading {level}"]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        style.paragraph_format.first_line_indent = Cm(0)
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE

    footer = section.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_page_number(fp)


def p(doc, text, *, indent=True, bold=False, center=False, size=12, space_after=6):
    para = doc.add_paragraph()
    para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    para.paragraph_format.space_after = Pt(space_after)
    para.paragraph_format.first_line_indent = Cm(0.75) if indent else Cm(0)
    if center:
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = para.add_run(text)
    set_run_font(run, size=size, bold=bold)
    return para


def caption(doc, text):
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.first_line_indent = Cm(0)
    para.paragraph_format.space_before = Pt(6)
    para.paragraph_format.space_after = Pt(12)
    run = para.add_run(text)
    set_run_font(run, size=11, italic=True)
    return para


def add_figure(doc, path: Path, caption_text: str, width_cm: float = 15.5):
    if not path.exists():
        p(doc, f"[Missing figure: {path.name}]", indent=False, bold=True)
        return
    para = doc.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.first_line_indent = Cm(0)
    run = para.add_run()
    run.add_picture(str(path), width=Cm(width_cm))
    caption(doc, caption_text)


def add_table(doc, headers, rows, col_widths=None):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ""
        para = hdr[i].paragraphs[0]
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = para.add_run(h)
        set_run_font(run, size=11, bold=True)
    for r_i, row in enumerate(rows):
        for c_i, val in enumerate(row):
            cell = table.rows[r_i + 1].cells[c_i]
            cell.text = ""
            para = cell.paragraphs[0]
            run = para.add_run(str(val))
            set_run_font(run, size=11)
    if col_widths:
        for row in table.rows:
            for idx, w in enumerate(col_widths):
                row.cells[idx].width = Cm(w)
    doc.add_paragraph()


def build() -> Path:
    doc = Document()
    style_doc(doc)

    # -------- Cover --------
    for _ in range(2):
        doc.add_paragraph()
    p(doc, "TSINGHUA UNIVERSITY", center=True, indent=False, bold=True, size=16, space_after=4)
    p(
        doc,
        "Innovation and Entrepreneurship Development Education (IEDE)",
        center=True,
        indent=False,
        size=12,
        space_after=4,
    )
    p(doc, "Autumn Semester 2026", center=True, indent=False, size=12, space_after=24)

    p(
        doc,
        "AI-Assisted Resilience for SDN-Enabled Intelligent Transportation Systems",
        center=True,
        indent=False,
        bold=True,
        size=18,
        space_after=8,
    )
    p(
        doc,
        "A Lightweight Framework for V2X Network Reliability, Edge Telemetry, and Safety-Aware Mitigation",
        center=True,
        indent=False,
        size=13,
        space_after=28,
    )
    p(doc, "Prototype Implementation & Ops-Console Demonstration Report", center=True, indent=False, bold=True, size=14, space_after=28)

    p(doc, "Author: Trinh Hoang Tu", center=True, indent=False, size=12, space_after=4)
    p(
        doc,
        "Project Advisor: Prof. Kris Singh (Visiting Professor, Tsinghua University; CEO, SRII)",
        center=True,
        indent=False,
        size=12,
        space_after=4,
    )
    p(doc, "Repository: sdn-its-resilience-ai (v0.2.0)", center=True, indent=False, size=12, space_after=4)
    p(doc, "Date: 30 July 2026", center=True, indent=False, size=12, space_after=4)

    doc.add_page_break()

    # -------- Abstract --------
    doc.add_heading("Abstract", level=1)
    p(
        doc,
        "Intelligent Transportation Systems (ITS) rely on Vehicle-to-Everything (V2X) connectivity, "
        "roadside units (RSUs), and software-defined networking (SDN) for programmable control. "
        "Centralized control improves agility but exposes the control plane to packet_in floods, "
        "low-rate OpenFlow trust (LOFT) probes, and false-positive blocking during legitimate rush-hour surges. "
        "This report presents a research prototype that couples flow-level telemetry, a hybrid TCN-GRU "
        "anomaly detector with adaptive threshold τ = μ_RE + kσ_RE, and a Detect → Contain → Rollback "
        "mitigation loop with cooldown. An ops console provides live inference, scenario replay, "
        "feature attribution, and OpenFlow rule visualization. Empirical smoke results indicate strong "
        "anomaly detection quality on synthetic and InSDN-shaped samples, sub-10 ms response on CPU for "
        "the decision path, and successful false-positive avoidance for rush-hour bursts.",
    )
    p(doc, "Keywords: SDN; ITS; V2X; TCN-GRU; anomaly detection; OpenFlow; resilience; cooldown rollback.", indent=False)

    doc.add_page_break()

    # -------- TOC --------
    doc.add_heading("Table of Contents", level=1)
    toc_hint = doc.add_paragraph()
    toc_hint.paragraph_format.first_line_indent = Cm(0)
    run = toc_hint.add_run(
        "Note: In Microsoft Word, right-click the table of contents below → Update Field → Update entire table."
    )
    set_run_font(run, size=10, italic=True)

    toc_para = doc.add_paragraph()
    toc_para.paragraph_format.first_line_indent = Cm(0)
    add_toc(toc_para)
    # Visible placeholder lines (replaced when TOC is updated in Word)
    for line in [
        "1. Introduction ........................................................ 3",
        "2. System Architecture ................................................. 4",
        "3. Prototype Implementation ............................................ 5",
        "4. Ops Console Demonstration (Annotated Screenshots) ................... 6",
        "5. Experimental Results ................................................ 10",
        "6. Discussion & Limitations ............................................ 11",
        "7. Conclusion & Future Work ............................................ 12",
        "References ............................................................. 13",
        "Appendix A. Figure Annotation Guide .................................... 14",
    ]:
        para = doc.add_paragraph(line)
        para.paragraph_format.first_line_indent = Cm(0)
        para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
        for r in para.runs:
            set_run_font(r, size=12)

    doc.add_page_break()

    # -------- 1 --------
    doc.add_heading("1. Introduction", level=1)
    doc.add_heading("1.1 Research Motivation", level=2)
    p(
        doc,
        "SDN enables centralized, programmable forwarding for ITS edge fabrics connecting vehicles, "
        "RSUs, and aggregation switches. However, a compromised or overloaded controller becomes a "
        "single point of failure. Permanent blacklists further risk disrupting legitimate rush-hour "
        "traffic. This project investigates a transportation-first resilience loop that detects "
        "anomalies from non-DPI flow counters and applies temporary, cooldown-bounded containment.",
    )
    doc.add_heading("1.2 Research Questions", level=2)
    for rq in [
        "RQ1: How can SDN telemetry detect abnormal ITS behavior without deep packet inspection?",
        "RQ2: Can a lightweight time-series model meet edge latency constraints while detecting floods and faults?",
        "RQ3: How can closed-loop mitigation reduce controller overload without false-positive service disruption?",
        "RQ4: How should public datasets (e.g., InSDN) be shaped with simulated VANET scenarios for evaluation?",
    ]:
        para = doc.add_paragraph(rq, style="List Number")
        para.paragraph_format.first_line_indent = Cm(0)
        for r in para.runs:
            set_run_font(r, size=12)

    doc.add_heading("1.3 Contributions of This Prototype", level=2)
    for item in [
        "Research-oriented package layout (src/sdn_its) with train/evaluate/quantize/infer modules.",
        "Detect → Contain → Rollback mitigation with policy matrix for DDoS, LOFT, RSU fault, and rush-hour.",
        "Ops console with live TCN-GRU scoring, online τ adaptation, scenario replay, and settings sync.",
        "InSDN download/prepare pipeline and ablation/benchmark tooling for paper-oriented evidence.",
    ]:
        para = doc.add_paragraph(item, style="List Bullet")
        para.paragraph_format.first_line_indent = Cm(0)
        for r in para.runs:
            set_run_font(r, size=12)

    # -------- 2 --------
    doc.add_heading("2. System Architecture", level=1)
    p(
        doc,
        "The end-to-end pipeline is: Vehicle → RSU → OVS → Controller → Telemetry → AI Model → "
        "Mitigation → Rollback. Telemetry windows of length L=10 with 10 flow-level features feed a "
        "hybrid TCN-GRU autoencoder/classifier. Reconstruction error (RE) is compared against an "
        "adaptive threshold τ. Policy actions push temporary OpenFlow rules (DROP / RATE_LIMIT / "
        "THROTTLE_PACKET_IN) that auto-expire via a cooldown daemon.",
    )
    p(
        doc,
        "The demonstration path uses a Flask REST controller with a mock OpenFlow push channel suitable "
        "for Windows development. A Ryu/Mininet adapter and lab checklist are reserved for Linux deployments.",
        indent=True,
    )

    # -------- 3 --------
    doc.add_heading("3. Prototype Implementation", level=1)
    doc.add_heading("3.1 Software Stack", level=2)
    add_table(
        doc,
        ["Layer", "Technology"],
        [
            ["AI model", "PyTorch TCN-GRU (FP32 + dynamic INT8 quantization)"],
            ["API / controller", "Flask REST (/api/v1/*)"],
            ["Dashboard", "Single-page ops console (IBM Plex), Chart.js"],
            ["Datasets", "Synthetic + InSDN (Hugging Face) → canonical CSV"],
            ["Experiments", "YAML-driven exp01–exp05 + ablation sweep"],
            ["Export", "ONNX / TorchScript, Markdown reports"],
        ],
        col_widths=[4.5, 11.0],
    )

    doc.add_heading("3.2 Mitigation Policy Matrix", level=2)
    add_table(
        doc,
        ["Alert type", "Action", "Notes"],
        [
            ["DDoS_PacketIn_Flood", "DROP + THROTTLE_PACKET_IN", "Cooldown-bounded containment"],
            ["LOFT_LowRate_Attack", "RATE_LIMIT", "Preserve partial connectivity"],
            ["RushHour_Burst", "IGNORE (no block)", "False-positive avoidance"],
            ["RSU_Link_Fault", "Reroute stub", "Lab hook for path failover"],
        ],
        col_widths=[5.0, 5.5, 5.0],
    )

    # -------- 4 Screenshots --------
    doc.add_heading("4. Ops Console Demonstration (Annotated Screenshots)", level=1)
    p(
        doc,
        "This section documents the working prototype UI captured on 30 July 2026 at "
        "http://127.0.0.1:8080/. Each figure is annotated for direct insertion into the IEDE dossier.",
    )

    doc.add_heading("4.1 Nominal Monitoring View", level=2)
    add_figure(
        doc,
        IMG / "fig01_console_overview.png",
        "Figure 1. Ops console overview under nominal monitoring. "
        "Annotation highlights: (A) pipeline strip Vehicle→…→Rollback; "
        "(B) live KPI cards (alerts, active rules, FP avoided, RE, τ, uptime); "
        "(C) RE vs τ chart with scenario injectors; "
        "(D) feature attribution bars; "
        "(E) ITS edge topology (s1 core, s2–s4 RSU, CTRL+AI); "
        "(F) event stream and rollback history.",
    )
    p(
        doc,
        "How to read Figure 1: When Active rules = 0, forwarding is nominal. The RE curve should remain "
        "below τ for normal telemetry. Feature attribution indicates which counters currently dominate RE.",
    )

    doc.add_heading("4.2 DDoS PacketIn Flood → Containment", level=2)
    add_figure(
        doc,
        IMG / "fig02_ddos_containment.png",
        "Figure 2. Active containment after Attack A (PacketIn flood / DDoS). "
        "Annotation highlights: (A) pipeline stages AI Model + Mitigation lit (detect/contain); "
        "(B) RE spike relative to τ; "
        "(C) Active OpenFlow rules table showing DROP on attacker src_ip and GLOBAL THROTTLE_PACKET_IN; "
        "(D) s1 CORE node highlighted as containment focus; "
        "(E) event stream line Policy · CONTAINED_DROP_AND_THROTTLE with policy latency in milliseconds.",
    )
    p(
        doc,
        "Interpretation: The policy engine applies temporary rules rather than permanent blacklists. "
        "Cooldown progress bars visualize remaining TTL until automatic rollback.",
    )

    doc.add_heading("4.3 Settings Drawer (Runtime Configuration)", level=2)
    add_figure(
        doc,
        IMG / "fig03_settings_drawer.png",
        "Figure 3. Settings drawer synchronized with POST /api/v1/settings. "
        "Annotation highlights: (A) console controls — poll interval, τ, DDoS/LOFT cooldown; "
        "(B) online k·σ, dataset mode, role mode (Research/Demo/Debug), theme; "
        "(C) online τ adapt and compare mode toggles; "
        "(D) optional webhook URL; "
        "(E) read-only controller/model metadata (version, seq_len, features, train CSV).",
        width_cm=14.0,
    )

    doc.add_heading("4.4 Rush-Hour Preserve (False-Positive Avoidance)", level=2)
    add_figure(
        doc,
        IMG / "fig04_rush_hour_preserve.png",
        "Figure 4. Benign rush-hour surge scenario. "
        "Annotation highlights: (A) scenario button Benign / Rush-hour surge; "
        "(B) policy result IGNORED_BENIGN_RUSH_HOUR (no DROP rules); "
        "(C) FP avoided counter increments; "
        "(D) compare-mode banner reminding reviewers to contrast permanent-block baselines vs cooldown rollback.",
    )
    p(
        doc,
        "This scenario demonstrates RQ3: legitimate transportation load must not be treated as an attack. "
        "The console intentionally separates rush-hour bursts from DDoS floods.",
    )

    doc.add_heading("4.5 Scenario Replay Timeline", level=2)
    add_figure(
        doc,
        IMG / "fig05_replay_timeline.png",
        "Figure 5. Scenario replay (DDoS cycle) mid-run. "
        "Annotation highlights: (A) replay script selector + Play/Stop; "
        "(B) progress track under the toolbar; "
        "(C) synchronized RE chart and event stream updates; "
        "(D) inference latency percentiles (p50 / p95) for edge-readiness discussion.",
    )

    doc.add_heading("4.6 Topology & Event Stream Detail", level=2)
    add_figure(
        doc,
        IMG / "fig06_topology_events.png",
        "Figure 6. Lower console region focusing on ITS edge topology and the OpenFlow/policy event stream. "
        "Annotation highlights: (A) vehicle groups → s1 CORE → RSU leaves s2/s3/s4 → CTRL+AI; "
        "(B) containment highlight on s1 when rules are active; "
        "(C) chronological policy/rollback messages used as audit evidence in demos.",
    )

    # -------- 5 Results --------
    doc.add_heading("5. Experimental Results", level=1)
    doc.add_heading("5.1 Detection Quality (Prototype Benchmark)", level=2)
    p(
        doc,
        "On the refined synthetic evaluation used in the IEDE summary, the TCN-GRU detector achieved "
        "the following binary anomaly metrics (threshold τ = μ_RE + 2σ_RE):",
    )
    add_table(
        doc,
        ["Metric", "Value"],
        [
            ["Anomaly accuracy", "97.33%"],
            ["Precision", "93.85%"],
            ["Recall", "100.00%"],
            ["F1-score", "96.83%"],
            ["Multi-class accuracy (5 classes)", "100.00% (synthetic separation)"],
        ],
        col_widths=[8.0, 7.0],
    )

    doc.add_heading("5.2 Latency & Footprint", level=2)
    add_table(
        doc,
        ["Metric", "FP32 (approx.)", "Notes"],
        [
            ["t_inference", "~0.19–3 ms", "CPU; batch=1"],
            ["IDL", "~6.4 ms", "collect + preprocess + score"],
            ["T_response", "~9.9 ms", "includes policy/API path"],
            ["Model size", "~0.98 MB / 0.17 MB INT8", "checkpoint vs quantized"],
        ],
        col_widths=[4.5, 5.0, 6.0],
    )
    p(
        doc,
        "Console-side inference percentiles (p50/p95) shown in Figures 1–5 provide complementary "
        "evidence for demo runs and may vary with host load.",
    )

    doc.add_heading("5.3 Ablation Snapshot", level=2)
    p(
        doc,
        "An architecture × k·σ ablation utility (scripts/run_sweep.py) compares RNN/LSTM/GRU baselines "
        "against TCN-GRU and reports rush-hour false-block rate as a transportation-continuity proxy. "
        "Detailed numeric tables are stored under results/ablation/ for paper drafting.",
    )

    # -------- 6 --------
    doc.add_heading("6. Discussion & Limitations", level=1)
    for item in [
        "OpenFlow actuation in the Windows demo path is mocked; production validation requires Ryu + Mininet-WiFi on Linux/WSL2.",
        "Live demo windows are standardized per window for readable RE scales; training uses a fitted StandardScaler on canonical CSVs.",
        "DRL orchestration (exp05) includes a replay-buffer training loop but is not yet the primary production policy.",
        "Public InSDN labels approximate ITS control-plane stressors; SUMO mobility coupling remains future work.",
    ]:
        para = doc.add_paragraph(item, style="List Bullet")
        para.paragraph_format.first_line_indent = Cm(0)
        for r in para.runs:
            set_run_font(r, size=12)

    # -------- 7 --------
    doc.add_heading("7. Conclusion & Future Work", level=1)
    p(
        doc,
        "The SDN-ITS Resilience prototype demonstrates a closed-loop path from flow telemetry to "
        "temporary OpenFlow containment with automatic rollback, packaged as a recruiter- and "
        "paper-oriented research framework. Immediate next steps are (i) real Ryu/Mininet experiments, "
        "(ii) multi-seed paper tables, and (iii) tighter coupling of online τ calibration to the trained scaler.",
    )

    doc.add_heading("References", level=1)
    refs = [
        "McKeown, N. et al. OpenFlow: Enabling innovation in campus networks. ACM SIGCOMM CCR, 2008.",
        "Bai, S., Kolter, J. Z., & Koltun, V. An Empirical Evaluation of Generic Convolutional and Recurrent Networks for Sequence Modeling. arXiv:1803.01271, 2018.",
        "Elsayed, M. S., Le-Khac, N.-A., & Jurcut, A. D. InSDN: A Novel SDN Intrusion Dataset. IEEE Access, 2020.",
        "Sharafaldin, I., Lashkari, A. H., & Ghorbani, A. A. Toward Generating a New Intrusion Detection Dataset and Intrusion Traffic Characterization (CSE-CIC-IDS2018).",
        "Trinh, H. T. AI-Assisted Resilience for SDN-Enabled ITS — IEDE Research Proposal Summary. Tsinghua University, 2026.",
    ]
    for i, ref in enumerate(refs, 1):
        para = doc.add_paragraph(f"[{i}] {ref}")
        para.paragraph_format.first_line_indent = Cm(0)
        para.paragraph_format.left_indent = Cm(0.75)
        para.paragraph_format.first_line_indent = Cm(-0.75)
        for r in para.runs:
            set_run_font(r, size=11)

    doc.add_page_break()
    doc.add_heading("Appendix A. Figure Annotation Guide (for DOCX editors)", level=1)
    p(
        doc,
        "Use these short labels when adding callouts/arrows in PowerPoint or Word drawing tools:",
        indent=False,
    )
    add_table(
        doc,
        ["Figure", "Suggested callout labels"],
        [
            ["Fig. 1", "A Pipeline · B KPIs · C RE chart · D Attribution · E Topology · F Events"],
            ["Fig. 2", "A Detect/Contain stages · B RE>τ · C DROP+THROTTLE rules · D s1 highlight · E Policy latency"],
            ["Fig. 3", "A Poll/τ/cooldown · B Role/theme · C Online τ · D Webhook · E Model metadata"],
            ["Fig. 4", "A Rush-hour button · B IGNORE policy · C FP avoided · D Compare banner"],
            ["Fig. 5", "A Replay controls · B Progress track · C Synced RE/events · D Latency p50/p95"],
            ["Fig. 6", "A Topology path · B Containment highlight · C Audit event stream"],
        ],
        col_widths=[3.0, 12.5],
    )
    p(
        doc,
        "Source PNG files: assets/images/report/fig01_*.png … fig06_*.png. "
        "Re-capture anytime with: python scripts/capture_report_screens.py",
        indent=False,
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUT))
    return OUT


if __name__ == "__main__":
    path = build()
    print(f"[+] Wrote {path}")
    print(f"    size={path.stat().st_size} bytes")
