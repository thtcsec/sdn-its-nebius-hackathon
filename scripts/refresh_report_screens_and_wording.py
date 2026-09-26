#!/usr/bin/env python
"""Refresh landscape screenshots, rebind into DOCX without stretch, fix captions/wording."""

from __future__ import annotations

import re
import shutil
import subprocess
import sys
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
DOC = ROOT / "Trinh Hoang Tu_3rd Version.docx"
OUT_DOCS = ROOT / "docs" / "IEDE_Project_Report_Prototype.docx"
IMG = ROOT / "assets" / "images" / "report"
ARCH = ROOT / "architecture.png"

FIG_SRC = {
    1: IMG / "fig01_console_overview.png",
    2: IMG / "fig02b_re_chart_ddos.png",
    3: IMG / "fig03_settings_drawer.png",
    4: IMG / "fig04_rush_hour_preserve.png",
    5: IMG / "fig05_replay_timeline.png",
    6: IMG / "fig06_topology_events.png",
    7: ARCH,
    8: IMG / "fig07_camera_tracks.png",
    9: IMG / "fig08_fusion_camera.png",
}

MAX_WIDTH_CM = 15.5
EMU_PER_CM = 360000

CAPTIONS = {
    1: (
        "Figure 1. Ops console overview (3 August 2026). "
        "Highlights: (A) evidence strip — MODEL / OF·mock / NET·NO EVIDENCE honesty; "
        "(B) pipeline stages Vehicle→…→Rollback; "
        "(C) KPI cards (policy decisions, in-memory rules, rush-hour preserved, RE, τ, uptime); "
        "(D) RE vs τ chart and network inject buttons (PacketIn ≠ camera); "
        "(E) intended mitigation rules panel (in-memory unless NET evidence)."
    ),
    2: (
        "Figure 2. PacketIn flood inject with RE chart and policy intents. "
        "Highlights: (A) RE rise under Attack A; "
        "(B) network inject note separating PacketIn from camera inputs; "
        "(C) intended DROP/THROTTLE or rate-limit rows in the rules panel; "
        "(D) Windows demo path remains mock OpenFlow unless Experiment A verifies the switch."
    ),
    3: (
        "Figure 3. Settings drawer (POST /api/v1/settings). "
        "Highlights: (A) poll interval, τ, DDoS/LOFT cooldown; "
        "(B) Lab mode / scenario label override (demo honesty); "
        "(C) online τ adapt (μ+kσ), role/theme; "
        "(D) ambient RE simulation toggle — leave off for reviews."
    ),
    4: (
        "Figure 4. Rush-hour network label scenario. "
        "Highlights: (A) Benign / Rush-hour surge inject; "
        "(B) GT may show RushHour_Burst while PRED can still flag flood — use MATCH strip; "
        "(C) policy should avoid hard DROP for legitimate surge when fusion preserves; "
        "(D) compare-mode reminder vs permanent-block baselines."
    ),
    5: (
        "Figure 5. Scenario replay (DDoS cycle). "
        "Highlights: (A) Replay · DDoS cycle + Play/Stop; "
        "(B) synchronized RE chart during contain window; "
        "(C) policy latency p50/p95; "
        "(D) feature bars (byte_rate, packet_in_count, …) for the active window."
    ),
    6: (
        "Figure 6. ITS edge topology and event stream. "
        "Highlights: (A) veh groups → s1 CORE → RSU s2/s3/s4 → CTRL+AI; "
        "(B) s1 containment highlight when rules are active; "
        "(C) chronological CONTAINED_DROP_AND_THROTTLE / rollback audit lines."
    ),
    7: (
        "Figure 7. Prototype software architecture: parallel camera context + SDN telemetry → "
        "TCN-GRU → multimodal fusion → policy → temporary OpenFlow (lab) / in-memory rules (Windows demo)."
    ),
    8: (
        "Figure 8. Camera context panel. "
        "Labels: Tracks (= visible_vehicle_tracks, not veh/h), q≈ queue_length_proxy_m, "
        "rel (= context_reliability, not calibrated confidence), source=synthetic or detector_tracker_v1. "
        "Stream quality flags (stream/frame/scene/analytics) are shown separately from track counts."
    ),
    9: (
        "Figure 9. Fusion decision (network × physical). "
        "PacketIn network inject is independent of Cam · * physical overrides. "
        "Outputs recommended_action + reason_codes (e.g., preserve_rush_hour vs temporary_drop)."
    ),
}

BODY_REPLACEMENTS = [
    (
        "This section documents the working prototype UI at http://127.0.0.1:8080/. Figures 1–6 and 8–9 were re-captured on 2 August 2026 from the current evidence console (Tracks / context_reliability labels, Cam · synth controls, NET · NO EVIDENCE honesty strip).",
        "This section documents the working prototype UI at http://127.0.0.1:8080/. "
        "Figures 1–6 and 8–9 were re-captured on 3 August 2026 in landscape (1440×900 @2× → 2880×1800). "
        "The evidence strip explicitly shows OF·mock and NET·NO EVIDENCE on Windows demos; "
        "camera cards use Tracks / rel rather than fake confidence scores.",
    ),
    (
        "How to read Figure 1: When Active rules = 0, forwarding is nominal. The RE curve should remain below τ for normal telemetry. Feature attribution indicates which counters currently dominate RE.",
        "How to read Figure 1: When pending/in-memory rules = 0, forwarding is nominal. "
        "The RE curve should remain below τ for calm windows. "
        "NET·NO EVIDENCE means rule rows are policy intents, not switch-verified OpenFlow installs.",
    ),
    (
        "Interpretation: The policy engine applies temporary rules rather than permanent blacklists. Cooldown progress bars visualize remaining TTL until automatic rollback.",
        "Interpretation: The policy engine prefers temporary, cooldown-bounded rules rather than permanent blacklists. "
        "On Windows mock OF, treat the rules table as intended mitigation; Experiment A provides switch dump evidence.",
    ),
]


def set_run_font(run, size=11, bold=False, italic=False):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic


def fill_para(para, text, *, size=11, italic=True, bold=False):
    for r in list(para.runs):
        r._element.getparent().remove(r._element)
    run = para.add_run(text)
    set_run_font(run, size=size, bold=bold, italic=italic)
    return para


def unpack(docx: Path, tmp: Path) -> None:
    if tmp.exists():
        shutil.rmtree(tmp)
    tmp.mkdir(parents=True)
    with zipfile.ZipFile(docx) as zf:
        zf.extractall(tmp)


def pack(tmp: Path, docx: Path) -> None:
    if docx.exists():
        docx.unlink()
    with zipfile.ZipFile(docx, "w", compression=zipfile.ZIP_DEFLATED) as zf:
        for f in tmp.rglob("*"):
            if f.is_file():
                zf.write(f, f.relative_to(tmp).as_posix())


def target_emus(path: Path) -> tuple[int, int]:
    im = Image.open(path)
    w, h = im.size
    cx = int(MAX_WIDTH_CM * EMU_PER_CM)
    cy = int(cx * (h / w))
    return cx, cy


def replace_nearest_extent(xml: str, caption_pos: int, cx: int, cy: int) -> str:
    window_start = max(0, caption_pos - 12000)
    head, mid, tail = xml[:window_start], xml[window_start:caption_pos], xml[caption_pos:]

    def _sub_last(fragment: str, tag: str) -> str:
        pat = re.compile(rf"(<{tag}\b[^>]*\bcx=\")\d+(\"\s+cy=\")\d+(\")")
        matches = list(pat.finditer(fragment))
        if not matches:
            return fragment
        m = matches[-1]
        return (
            fragment[: m.start()]
            + m.group(1)
            + str(cx)
            + m.group(2)
            + str(cy)
            + m.group(3)
            + fragment[m.end() :]
        )

    mid = _sub_last(mid, "wp:extent")
    mid = _sub_last(mid, "a:ext")
    return head + mid + tail


def rebind_images() -> None:
    for _n, p in FIG_SRC.items():
        if not p.exists():
            raise SystemExit(f"missing {p}")
    tmp = ROOT / "tmp_docx_rebind"
    unpack(DOC, tmp)
    xml_path = tmp / "word" / "document.xml"
    xml = xml_path.read_text(encoding="utf-8")
    rels = (tmp / "word" / "_rels" / "document.xml.rels").read_text(encoding="utf-8")
    rid_to = dict(re.findall(r'Id="(rId\d+)"[^>]*Target="media/([^"]+)"', rels))

    for n, src in sorted(FIG_SRC.items()):
        m = re.search(rf"Figure\s+{n}\.", xml)
        if not m:
            print(f"  skip Figure {n}: no caption")
            continue
        window = xml[max(0, m.start() - 12000) : m.start()]
        embeds = re.findall(r'r:embed="(rId\d+)"', window)
        if not embeds:
            print(f"  skip Figure {n}: no blip")
            continue
        rid = embeds[-1]
        media_name = rid_to[rid]
        dest = tmp / "word" / "media" / media_name
        if dest.suffix.lower() in {".jpg", ".jpeg"}:
            Image.open(src).convert("RGB").save(dest, format="JPEG", quality=92)
        else:
            shutil.copyfile(src, dest)
        cx, cy = target_emus(src)
        xml = replace_nearest_extent(xml, m.start(), cx, cy)
        im = Image.open(src)
        print(f"  Figure {n}: {media_name} <- {src.name} {im.size[0]}x{im.size[1]}")

    xml_path.write_text(xml, encoding="utf-8")
    pack(tmp, DOC)


def fix_wording() -> None:
    doc = Document(str(DOC))
    for p in doc.paragraphs:
        t = p.text or ""
        # exact body replacements
        for old, new in BODY_REPLACEMENTS:
            if t.strip() == old.strip() or t.startswith(old[:80]):
                # replace if starts with old prefix (robust)
                if old[:60] in t or t.strip().startswith(old[:40]):
                    fill_para(p, new, size=12, italic=False)
                    break
        # captions
        for n, caption in CAPTIONS.items():
            if t.startswith(f"Figure {n}."):
                fill_para(p, caption, size=11, italic=True)
                break
        # date
        if t.strip() == "August, 2026":
            fill_para(p, "August, 2026", size=12, italic=False)
        # architecture prose already OK

    # Appendix annotation table — last table if Figure/Suggested
    if doc.tables:
        t_last = doc.tables[-1]
        if (t_last.rows[0].cells[0].text or "").lower().startswith("figure"):
            updates = {
                "Fig. 1": "A Evidence strip · B Pipeline · C KPIs · D RE chart · E Rules panel",
                "Fig. 2": "A RE rise · B PacketIn≠camera note · C Intended rules · D Mock OF caveat",
                "Fig. 3": "A Poll/τ/cooldown · B Lab/override · C Online τ · D Ambient RE off for review",
                "Fig. 4": "A Rush-hour button · B GT/PRED/MATCH · C Preserve/no hard DROP · D Compare banner",
                "Fig. 5": "A Replay controls · B RE during contain · C p50/p95 · D Feature bars",
                "Fig. 6": "A Topology path · B s1 highlight · C Event audit stream",
                "Fig. 7": "A Parallel cameras · B Telemetry/AI · C Fusion · D Policy · E Lab vs mock OF · F Console",
                "Fig. 8": "A Cam modes · B Tracks/rel · C Stream quality · D synthetic/detector source",
                "Fig. 9": "A Network evidence · B Physical tracks · C recommended_action · D explain",
            }
            for row in t_last.rows[1:]:
                key = (row.cells[0].text or "").strip()
                if key in updates:
                    row.cells[1].text = ""
                    run = row.cells[1].paragraphs[0].add_run(updates[key])
                    set_run_font(run, size=11)

            # ensure Fig 8/9 rows exist
            existing = {(r.cells[0].text or "").strip() for r in t_last.rows}
            for key, val in (("Fig. 8", updates["Fig. 8"]), ("Fig. 9", updates["Fig. 9"])):
                if key not in existing:
                    row = t_last.add_row()
                    row.cells[0].text = ""
                    row.cells[1].text = ""
                    r0 = row.cells[0].paragraphs[0].add_run(key)
                    set_run_font(r0, size=11)
                    r1 = row.cells[1].paragraphs[0].add_run(val)
                    set_run_font(r1, size=11)

    # Soft-fix §4.7/4.8 body if still OK — leave as is if already honest
    doc.save(str(DOC))
    print("wording/captions updated")


def main() -> None:
    # 1) refresh screenshots
    rc = subprocess.call([sys.executable, str(ROOT / "scripts" / "capture_report_screens_landscape.py")])
    if rc != 0:
        raise SystemExit(f"capture failed: {rc}")
    # fallback if fig02b missing
    if not FIG_SRC[2].exists():
        FIG_SRC[2] = IMG / "fig02_ddos_containment.png"

    # 2) rebind images + extents
    rebind_images()

    # 3) wording
    fix_wording()

    shutil.copyfile(DOC, OUT_DOCS)
    d = Document(str(DOC))
    imgs = sum(1 for r in d.part.rels.values() if "image" in r.reltype)
    print(f"OK {DOC.name}: paras={len(d.paragraphs)} tables={len(d.tables)} images={imgs}")
    print(f"copied -> {OUT_DOCS}")


if __name__ == "__main__":
    main()
