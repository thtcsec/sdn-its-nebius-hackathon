#!/usr/bin/env python
"""Bind live-detect camera figures into DOCX and fix captions."""

from __future__ import annotations

import re
import shutil
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
DOC = ROOT / "Trinh Hoang Tu_3rd Version.docx"
OUT = ROOT / "docs" / "IEDE_Project_Report_Prototype.docx"
IMG = ROOT / "assets" / "images" / "report"

FIG = {
    8: IMG / "fig07_camera_tracks.png",
    9: IMG / "fig08_fusion_camera.png",
}
CAPS = {
    8: (
        "Figure 8. Live RTSP camera frames with YOLO object detection overlays "
        "(cars / motorcycles / persons). Captured while streams were reachable; "
        "boxes from detector_tracker_v1. This is object detection — not an "
        "unusual-event or accident classifier."
    ),
    9: (
        "Figure 9. Close-up of an annotated live camera frame (cam1). "
        "Boxes show vehicles and riders in-view. Physical context for fusion uses "
        "visible_vehicle_tracks and context_reliability — not a claim of unusual-behavior AI."
    ),
}
BODY_47 = (
    "Figure 8 shows live RTSP frames with YOLO boxes (cars/motorcycles/persons) from "
    "detector_tracker_v1. These are object detections used as physical context "
    "(visible_vehicle_tracks / context_reliability) — not an unusual-event, accident, "
    "or anomaly-behavior detector. Current DDNS may be unreachable; frames were captured "
    "when streams were online."
)

EMU = 360000
MAXW = 15.5


def font(run, size=11, italic=True) -> None:
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.italic = italic


def main() -> None:
    for p in FIG.values():
        if not p.exists():
            raise SystemExit(f"missing {p}")

    tmp = ROOT / "tmp_docx_camfix"
    if tmp.exists():
        shutil.rmtree(tmp)
    tmp.mkdir()
    with zipfile.ZipFile(DOC) as zf:
        zf.extractall(tmp)

    xml_path = tmp / "word" / "document.xml"
    xml = xml_path.read_text(encoding="utf-8")
    rels = (tmp / "word" / "_rels" / "document.xml.rels").read_text(encoding="utf-8")
    rid_to = dict(re.findall(r'Id="(rId\d+)"[^>]*Target="media/([^"]+)"', rels))

    for n, src in FIG.items():
        m = re.search(rf"Figure\s+{n}\.", xml)
        if not m:
            raise SystemExit(f"caption Figure {n} missing")
        window = xml[max(0, m.start() - 12000) : m.start()]
        embeds = re.findall(r'r:embed="(rId\d+)"', window)
        rid = embeds[-1]
        media = rid_to[rid]
        dest = tmp / "word" / "media" / media
        shutil.copyfile(src, dest)
        im = Image.open(src)
        w, h = im.size
        cx = int(MAXW * EMU)
        cy = int(cx * (h / w))
        frag_start = max(0, m.start() - 12000)
        head, mid, tail = xml[:frag_start], xml[frag_start : m.start()], xml[m.start() :]
        for tag in ("wp:extent", "a:ext"):
            pat = re.compile(rf"(<{tag}\b[^>]*\bcx=\")\d+(\"\s+cy=\")\d+(\")")
            ms = list(pat.finditer(mid))
            if ms:
                mm = ms[-1]
                mid = (
                    mid[: mm.start()]
                    + mm.group(1)
                    + str(cx)
                    + mm.group(2)
                    + str(cy)
                    + mm.group(3)
                    + mid[mm.end() :]
                )
        xml = head + mid + tail
        print(f"bound Figure {n}: {media} {w}x{h}")

    xml_path.write_text(xml, encoding="utf-8")
    packed = ROOT / "Trinh Hoang Tu_3rd Version.camtmp.docx"
    if packed.exists():
        packed.unlink()
    with zipfile.ZipFile(packed, "w", compression=zipfile.ZIP_DEFLATED) as zf:
        for f in tmp.rglob("*"):
            if f.is_file():
                zf.write(f, f.relative_to(tmp).as_posix())
    shutil.move(packed, DOC)

    doc = Document(str(DOC))
    for p in doc.paragraphs:
        t = p.text or ""
        for n, cap in CAPS.items():
            if t.startswith(f"Figure {n}."):
                for r in list(p.runs):
                    r._element.getparent().remove(r._element)
                run = p.add_run(cap)
                font(run, 11, True)
                print(f"caption {n}")
        if t.startswith("The ops console camera cards expose") or t.startswith("Figure 8 shows live RTSP"):
            for r in list(p.runs):
                r._element.getparent().remove(r._element)
            run = p.add_run(BODY_47)
            font(run, 12, False)
            print("updated 4.7 body")
    doc.save(str(DOC))
    shutil.copyfile(DOC, OUT)
    print(f"wrote {DOC.name}")
    print(f"copied {OUT}")


if __name__ == "__main__":
    main()
